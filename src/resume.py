"""음성 파일과 짝이 되는 '이력서' 파일을 찾아 읽는다.

매칭 규칙: 음성 파일과 '같은 이름'(확장자만 다름)이거나, 음성 파일 이름으로
시작하는 이력서 파일을 input 폴더에서 찾는다.
  예) 정주호.m4a  ↔  정주호.pdf  /  정주호_이력서.pdf

읽기 방식:
  - PDF/이미지 : Gemini가 직접 읽도록 '파일'로 전달 (바이트 + mime)
  - docx/txt   : 텍스트로 추출해 전달
  - 그 외(.hwp 등): 미지원 → 전화 내용만 사용 (PDF로 변환 권장)
"""
from __future__ import annotations

import re
from pathlib import Path

from . import config

_IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


def first_token(name: str) -> str:
    """파일 이름에서 첫 단어(보통 사람 이름)를 뽑는다. 예: '정주원 이력서' → '정주원'."""
    parts = re.split(r"[\s_]+", name.strip())
    return parts[0] if parts and parts[0] else name.strip()


def find_resume(audio_path: Path) -> Path | None:
    """음성 파일과 짝이 되는 이력서 파일 경로를 찾는다. 없으면 None.

    매칭 우선순위:
      1) 같은 이름, 확장자만 다름        (정주원.m4a ↔ 정주원.pdf)
      2) 한쪽 이름이 다른 쪽으로 시작     (정주원.m4a ↔ 정주원_이력서.pdf)
      3) 첫 단어(이름)가 같음            (정주원 통화녹음.m4a ↔ 정주원 이력서.pdf)
    """
    # 1) 같은 이름, 확장자만 다른 경우
    for ext in config.RESUME_EXTS:
        cand = audio_path.with_suffix(ext)
        if cand.exists():
            return cand

    try:
        candidates = [
            p
            for p in sorted(audio_path.parent.iterdir())
            if p != audio_path and p.suffix.lower() in config.RESUME_EXTS
        ]
    except FileNotFoundError:
        return None

    a_stem = audio_path.stem

    # 2) 한쪽 이름이 다른 쪽 이름으로 시작
    for p in candidates:
        if p.stem.startswith(a_stem) or a_stem.startswith(p.stem):
            return p

    # 3) 첫 단어(이름)가 같음
    a_tok = first_token(a_stem)
    if a_tok:
        for p in candidates:
            if first_token(p.stem) == a_tok:
                return p

    return None


def load_resume(path: Path) -> dict:
    """이력서를 Gemini 입력 형태로 읽는다.

    반환:
      {"kind": "file", "data": bytes, "mime": str}  # PDF/이미지
      {"kind": "text", "text": str}                  # docx/txt
      {"kind": "unsupported", "ext": str}            # 그 외
    """
    ext = path.suffix.lower()

    if ext == ".pdf":
        return {"kind": "file", "data": path.read_bytes(), "mime": "application/pdf"}

    if ext in _IMAGE_MIME:
        return {"kind": "file", "data": path.read_bytes(), "mime": _IMAGE_MIME[ext]}

    if ext == ".txt":
        return {"kind": "text", "text": path.read_text(encoding="utf-8", errors="ignore")}

    if ext == ".docx":
        from docx import Document

        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return {"kind": "text", "text": "\n".join(lines)}

    return {"kind": "unsupported", "ext": ext}
