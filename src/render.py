"""추출 결과(CandidateComment)를 후보자 코멘트 .docx 문서로 생성한다.

업로드된 양식의 구조를 그대로 따라 새 문서를 만든다(빈칸 치환보다 안정적).
양식 구조:
    [직무명] '[성명]' 후보자 코멘트   (제목)
    [성명]
    경력 ...
    현재 연봉 ...
    희망 연봉 ...
    [내용]
    ...
    [상황]
    ...
    ◦ 각 재직 기업 이직 사유
    ...
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .extract import CandidateComment


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)


def render(comment: CandidateComment, output_path: Path) -> Path:
    doc = Document()

    # 제목
    job = comment.job_title or "[직무명]"
    name = comment.name or "[성명]"
    title = doc.add_paragraph()
    trun = title.add_run(f"[{job}] '{name}' 후보자 코멘트")
    trun.bold = True
    trun.font.size = Pt(13)

    # 기본 정보
    doc.add_paragraph(name)
    if comment.career:
        doc.add_paragraph(f"경력 {comment.career}.")
    if comment.current_salary:
        doc.add_paragraph(f"현재 연봉 {comment.current_salary}.")
    if comment.desired_salary:
        doc.add_paragraph(f"희망 연봉 {comment.desired_salary}.")

    # [내용]
    _heading(doc, "[내용]")
    if comment.skills:
        doc.add_paragraph(", ".join(comment.skills) + " 등")
    if comment.summary:
        doc.add_paragraph(comment.summary)
    if comment.desired_role:
        doc.add_paragraph(comment.desired_role)

    # [상황]
    _heading(doc, "[상황]")
    if comment.job_seeking_status:
        doc.add_paragraph(comment.job_seeking_status)

    if comment.histories:
        doc.add_paragraph("◦ 각 재직 기업 이직 사유")
        for h in comment.histories:
            parts = []
            if h.company:
                parts.append(f"‘{h.company}’")
            if h.period:
                parts.append(f"{h.period}간")
            if h.work:
                parts.append(h.work)
            line = " ".join(parts).strip()
            if h.reason:
                line = f"{line} {h.reason}".strip()
            if line:
                doc.add_paragraph(line)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
